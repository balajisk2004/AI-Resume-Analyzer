import PyPDF2
import docx
import re
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import spacy
import numpy as np
from datetime import datetime

# Load spaCy model
try:
    nlp = spacy.load('en_core_web_sm')
except:
    import subprocess
    subprocess.run(['python', '-m', 'spacy', 'download', 'en_core_web_sm'])
    nlp = spacy.load('en_core_web_sm')

class ResumeParser:
    COMMON_SKILLS = {
        'python', 'java', 'javascript', 'react', 'django', 'flask', 'spring boot',
        'html', 'css', 'sql', 'mongodb', 'postgresql', 'mysql', 'aws', 'azure',
        'docker', 'kubernetes', 'git', 'machine learning', 'deep learning',
        'tensorflow', 'pytorch', 'nlp', 'data analysis', 'pandas', 'numpy',
        'scikit-learn', 'rest api', 'graphql', 'node.js', 'express', 'angular',
        'vue.js', 'typescript', 'php', 'laravel', 'ruby', 'swift', 'kotlin',
        'android', 'ios', 'devops', 'ci/cd', 'jenkins', 'linux', 'agile', 'scrum',
        'c++', 'c#', '.net', 'go', 'rust', 'salesforce', 'sap', 'oracle', 'redis',
        'elasticsearch', 'hadoop', 'spark', 'tableau', 'power bi', 'excel', 'word',
        'photoshop', 'illustrator', 'figma', 'ui/ux', 'product management', 'project management'
    }
    
    @staticmethod
    def extract_text_from_pdf(file_path):
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e:
            print(f"PDF Error: {e}")
        return text
    
    @staticmethod
    def extract_text_from_docx(file_path):
        text = ""
        try:
            doc = docx.Document(file_path)
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text += cell.text + "\n"
        except Exception as e:
            print(f"DOCX Error: {e}")
        return text
    
    @staticmethod
    def extract_skills(text):
        text_lower = text.lower()
        found_skills = set()
        
        for skill in ResumeParser.COMMON_SKILLS:
            pattern = r'\b' + re.escape(skill) + r'\b'
            if re.search(pattern, text_lower):
                found_skills.add(skill)
        
        # Check for skill variations
        skill_variations = {
            'javascript': ['js', 'javascript', 'ecmascript'],
            'python': ['python', 'py', 'django', 'flask'],
            'react': ['reactjs', 'react js', 'react.js'],
            'machine learning': ['ml', 'machine learning', 'deep learning'],
        }
        
        for main_skill, variations in skill_variations.items():
            for var in variations:
                if var in text_lower:
                    found_skills.add(main_skill)
        
        return ', '.join(sorted(found_skills)) if found_skills else 'General Skills'
    
    @staticmethod
    def extract_education(text):
        education_keywords = [
            'bachelor', 'master', 'phd', 'b.sc', 'm.sc', 'b.tech', 'm.tech', 
            'b.e', 'm.e', 'degree', 'diploma', 'university', 'college', 
            'institute', 'school', 'academy', 'bca', 'mca', 'b.com', 'm.com',
            'b.a', 'm.a', 'b.b.a', 'm.b.a', 'ph.d', 'doctorate'
        ]
        
        sentences = re.split(r'[.!?\n]', text)
        education_lines = []
        
        for sentence in sentences:
            sentence_lower = sentence.lower()
            if any(keyword in sentence_lower for keyword in education_keywords):
                clean_sentence = sentence.strip()
                if 10 < len(clean_sentence) < 200:
                    education_lines.append(clean_sentence)
        
        return '\n'.join(education_lines[:5]) if education_lines else 'Education information not found'
    
    @staticmethod
    def extract_experience(text):
        experience_keywords = [
            'experience', 'worked', 'developed', 'managed', 'led', 'created', 
            'built', 'implemented', 'designed', 'architected', 'responsible for',
            'position', 'role', 'job', 'employment', 'work history', 'career',
            'senior', 'lead', 'principal', 'head of'
        ]
        
        lines = text.split('\n')
        experience_lines = []
        
        for line in lines:
            line_lower = line.lower()
            if any(keyword in line_lower for keyword in experience_keywords):
                clean_line = line.strip()
                if 20 < len(clean_line) < 300:
                    experience_lines.append(clean_line)
        
        # Extract years of experience
        year_pattern = r'(\d+)\+?\s*(?:years?|yrs?)'
        years_matches = re.findall(year_pattern, text.lower())
        if years_matches:
            max_years = max(int(y) for y in years_matches)
            experience_lines.append(f"Total Experience: {max_years}+ years")
        
        return '\n'.join(experience_lines[:10]) if experience_lines else 'Experience information not found'
    
    @staticmethod
    def extract_years_of_experience(text):
        """Extract years of experience from resume text"""
        year_patterns = [
            r'(\d+)\+?\s*(?:years?|yrs?)\s+of\s+experience',
            r'(\d+)\+?\s*(?:years?|yrs?)',
            r'experience\s+of\s+(\d+)\+?\s*(?:years?|yrs?)'
        ]
        
        for pattern in year_patterns:
            matches = re.findall(pattern, text.lower())
            if matches:
                return max(int(m) for m in matches)
        return 0
    
    @classmethod
    def parse_resume(cls, file_path):
        print(f"Parsing file: {file_path}")
        
        if file_path.endswith('.pdf'):
            text = cls.extract_text_from_pdf(file_path)
        elif file_path.endswith('.docx'):
            text = cls.extract_text_from_docx(file_path)
        else:
            raise ValueError("Unsupported file format")
        
        print(f"Extracted text length: {len(text)} characters")
        
        if not text or len(text.strip()) < 50:
            raise ValueError("Could not extract text from the file. Please ensure it's not encrypted or corrupted.")
        
        skills = cls.extract_skills(text)
        education = cls.extract_education(text)
        experience = cls.extract_experience(text)
        years_exp = cls.extract_years_of_experience(text)
        
        print(f"Extracted skills: {skills[:100]}...")
        print(f"Extracted years: {years_exp}")
        
        return {
            'full_text': text[:15000],
            'skills': skills,
            'education': education,
            'experience': experience,
            'years_experience': years_exp
        }

class ATS_Calculator:
    @staticmethod
    def calculate_ats_score(resume_text, job_description):
        """Calculate ATS score based on keyword matching and formatting"""
        score = 0
        explanations = []
        
        # Check for common sections (20 points)
        sections = ['education', 'experience', 'skills', 'summary', 'projects', 'certifications']
        found_sections = 0
        for section in sections:
            if re.search(r'\b' + section + r'\b', resume_text.lower()):
                found_sections += 1
        
        section_score = (found_sections / len(sections)) * 20
        score += section_score
        
        # Keyword matching (50 points)
        job_keywords = set(re.findall(r'\b[a-z]{3,}\b', job_description.lower()))
        resume_keywords = set(re.findall(r'\b[a-z]{3,}\b', resume_text.lower()))
        
        common_keywords = len(job_keywords.intersection(resume_keywords))
        if job_keywords:
            keyword_score = (common_keywords / len(job_keywords)) * 50
            score += keyword_score
        
        # Length and formatting (30 points)
        words = len(resume_text.split())
        if 400 <= words <= 800:
            score += 30
        elif 300 <= words < 400:
            score += 25
        elif 200 <= words < 300:
            score += 20
        elif words > 800:
            score += 15
        else:
            score += 5
        
        # Bonus for bullet points
        bullet_count = resume_text.count('•') + resume_text.count('-') + resume_text.count('*')
        if bullet_count > 20:
            score += 10
        elif bullet_count > 10:
            score += 5
        
        final_score = round(min(score, 100), 2)
        return final_score

class JobMatcher:
    def __init__(self):
        self.vectorizer = TfidfVectorizer(max_features=1000, stop_words='english')
    
    def calculate_enhanced_match_score(self, resume_skills, job_skills, resume_text, job_description, resume_experience_text, job_experience_required):
        """Enhanced matching algorithm considering skills, text similarity, and experience"""
        
        resume_skills_set = set(resume_skills)
        job_skills_set = set(job_skills)
        
        # 1. Skills Matching (50% weight)
        if job_skills_set and resume_skills_set:
            matched_skills = resume_skills_set.intersection(job_skills_set)
            missing_skills = job_skills_set - resume_skills_set
            skills_score = (len(matched_skills) / len(job_skills_set)) * 50
        else:
            matched_skills = set()
            missing_skills = set()
            skills_score = 0
        
        # 2. Text Similarity Matching (30% weight)
        try:
            if len(resume_text.strip()) > 100 and len(job_description.strip()) > 100:
                texts = [resume_text[:5000], job_description[:5000]]
                tfidf_matrix = self.vectorizer.fit_transform(texts)
                similarity = cosine_similarity(tfidf_matrix[0:1], tfidf_matrix[1:2])[0][0]
                text_score = similarity * 30
            else:
                text_score = 0
        except:
            text_score = 0
        
        # 3. Experience Matching (20% weight)
        exp_score = 0
        if job_experience_required > 0:
            # Extract years from resume experience text
            year_pattern = r'(\d+)\+?\s*(?:years?|yrs?)'
            experience_years = re.findall(year_pattern, resume_experience_text.lower())
            if experience_years:
                candidate_years = max(int(y) for y in experience_years)
                if candidate_years >= job_experience_required:
                    exp_score = 20
                elif candidate_years >= job_experience_required - 2:
                    exp_score = 15
                elif candidate_years >= job_experience_required - 4:
                    exp_score = 10
                else:
                    exp_score = 5
            else:
                exp_score = 10  # Default if no experience mentioned
        else:
            exp_score = 20  # No experience required
        
        total_score = skills_score + text_score + exp_score
        
        # Generate intelligent recommendations
        recommendations = []
        if missing_skills:
            recommendations.append(f" Add these skills to improve match: {', '.join(list(missing_skills)[:5])}")
        if exp_score < 15 and job_experience_required > 0:
            recommendations.append(f" Highlight your experience (need {job_experience_required}+ years)")
        if total_score < 50:
            recommendations.append(" Customize your resume for this specific role")
        if len(resume_text.split()) < 300:
            recommendations.append("📝 Add more details to your resume (aim for 300-800 words)")
        if not matched_skills:
            recommendations.append(" Highlight relevant skills and technologies in your resume")
        
        return {
            'match_score': round(total_score, 2),
            'matched_skills': ', '.join(sorted(matched_skills)) if matched_skills else 'None',
            'missing_skills': ', '.join(sorted(missing_skills)) if missing_skills else 'None',
            'recommendations': '\n'.join(recommendations) if recommendations else '✅ Your resume looks well-matched for this role!'
        }